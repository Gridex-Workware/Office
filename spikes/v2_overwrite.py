"""V2: 기존 Google 문서를 수정한 .docx/.pptx로 덮어쓰면 무엇이 유지되는지 확인.

원본은 건드리지 않는다. 원본의 사본을 만들어 그 사본으로 실험한다.

사용법: uv run python spikes/v2_overwrite.py <문서 또는 프레젠테이션 URL>
결과:   spikes/out/v2/report.md
"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import docx
import pptx
from _common import EXPORTS, FIELDS, OUT, credentials, drive_service, export, get_meta
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaFileUpload
from pptx.util import Inches

from officesync.content import extract_text


def snapshot(drive, file_id: str) -> dict:
    meta = drive.files().get(fileId=file_id, fields=FIELDS, supportsAllDrives=True).execute()
    perms = (
        drive.permissions()
        .list(fileId=file_id, fields="permissions(id,role,type)", supportsAllDrives=True)
        .execute()
    )
    meta["_permissions"] = sorted((p["id"], p["role"]) for p in perms.get("permissions", []))
    try:
        revisions = drive.revisions().list(fileId=file_id, fields="revisions(id)").execute()
        meta["_revisions"] = len(revisions.get("revisions", []))
    except HttpError:
        meta["_revisions"] = -1  # 확인 불가
    return meta


def add_marker(path: Path, marker: str) -> Path:
    edited = path.with_name(f"{path.stem}_edited{path.suffix}")
    if path.suffix == ".docx":
        document = docx.Document(str(path))
        document.add_paragraph(marker)
        document.save(edited)
    else:
        presentation = pptx.Presentation(str(path))
        slide = presentation.slides[0]
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(0.6))
        box.text_frame.text = marker
        presentation.save(edited)
    return edited


def upload(drive, file_id: str, path: Path, mime: str, google_mime: str) -> tuple[dict, str]:
    """덮어쓰기. 먼저 내용만 올려 보고, 실패하면 Google 형식을 명시해 다시 시도한다."""
    attempts = [("내용만 업로드", None), ("Google 형식 명시", {"mimeType": google_mime})]
    errors = []
    for label, body in attempts:
        media = MediaFileUpload(str(path), mimetype=mime, resumable=True)
        try:
            result = (
                drive.files()
                .update(
                    fileId=file_id,
                    body=body,
                    media_body=media,
                    fields=FIELDS,
                    supportsAllDrives=True,
                )
                .execute()
            )
            return result, label
        except HttpError as e:
            errors.append(f"{label}: {e}")
            print(f"실패 — {label}: {e}")
    raise SystemExit("두 방식 모두 실패했습니다:\n" + "\n".join(errors))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("url", help="Google 문서/프레젠테이션 주소 (원본은 바뀌지 않음)")
    args = parser.parse_args()

    creds = credentials()
    drive = drive_service(creds)
    out = OUT / "v2"

    original = get_meta(drive, args.url)
    mime, _ = EXPORTS[original["mimeType"]]
    copy = (
        drive.files()
        .copy(
            fileId=original["id"],
            body={"name": f"[officesync V2 테스트] {original['name']}"},
            fields="id",
            supportsAllDrives=True,
        )
        .execute()
    )
    print(f"사본 생성: {copy['id']}")

    before = snapshot(drive, copy["id"])
    exported = export(drive, creds, before, out)
    marker = f"officesync V2 marker {datetime.now():%Y%m%d-%H%M%S}"
    edited = add_marker(exported, marker)

    _, method = upload(drive, copy["id"], edited, mime, original["mimeType"])
    after = snapshot(drive, copy["id"])
    # Google 형식이 아닌 파일로 바뀌었다면 내보낼 수 없다 (형식 행에서 ❌로 드러남)
    marker_found = after["mimeType"] in EXPORTS and marker in extract_text(
        export(drive, creds, after, out, suffix="_after")
    )

    def same(key: str) -> str:
        return "✅ 유지" if before[key] == after[key] else "❌ 바뀜"

    rows = [
        ("파일 ID", before["id"], after["id"], same("id")),
        ("형식", before["mimeType"], after["mimeType"], same("mimeType")),
        ("제목", before["name"], after["name"], same("name")),
        ("링크", before["webViewLink"], after["webViewLink"], same("webViewLink")),
        (
            "공유 권한 수",
            len(before["_permissions"]),
            len(after["_permissions"]),
            same("_permissions"),
        ),
        (
            "version",
            before["version"],
            after["version"],
            "✅ 증가" if int(after["version"]) > int(before["version"]) else "❌ 그대로",
        ),
        (
            "버전 기록 수",
            before["_revisions"],
            after["_revisions"],
            "✅ 늘어남" if after["_revisions"] > before["_revisions"] else "⚠ 그대로",
        ),
        (
            "수정 내용 반영",
            "—",
            "마커 있음" if marker_found else "마커 없음",
            "✅" if marker_found else "❌",
        ),
    ]

    lines = [
        f"# V2 덮어쓰기 결과 ({datetime.now():%Y-%m-%d %H:%M})",
        "",
        f"- 원본: {original['name']} (변경 없음)",
        f"- 테스트 사본: {after['webViewLink']}",
        f"- 성공한 업로드 방식: **{method}**",
        "",
        "| 항목 | 덮어쓰기 전 | 덮어쓴 후 | 결과 |",
        "|---|---|---|---|",
        *(f"| {a} | {b} | {c} | {d} |" for a, b, c, d in rows),
        "",
        "## 브라우저에서 직접 확인할 것",
        "",
        "- [ ] 사본의 서식(글꼴, 표, 이미지)이 원본과 비슷한가",
        "- [ ] 파일 → 버전 기록에서 덮어쓰기 전 버전으로 되돌릴 수 있는가",
    ]
    report = out / "report.md"
    report.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    print(f"\n보고서: {report}")

    answer = input("\n확인이 끝났으면 테스트 사본을 휴지통으로 옮길까요? [y/N] ")
    if answer.strip().lower() == "y":
        drive.files().update(
            fileId=copy["id"], body={"trashed": True}, supportsAllDrives=True
        ).execute()
        print("휴지통으로 옮겼습니다.")


if __name__ == "__main__":
    main()
