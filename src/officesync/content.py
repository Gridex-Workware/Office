"""docx/pptx 내용 지문.

편집기는 수정 없이 저장만 해도 파일 바이트를 바꾼다 (zip 시각, 속성 순서, 런 분할 등).
그래서 '수정됨' 판정은 바이트 대신 텍스트와 서식으로 만든 지문으로 한다.
"""

from __future__ import annotations

import hashlib
import json
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from pathlib import Path

import docx
import pptx
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass(frozen=True)
class Fingerprint:
    text: str  # 텍스트만
    format: str  # 텍스트 + 서식 + 이미지 + (슬라이드) 도형 위치


def fingerprint(path: Path) -> Fingerprint:
    text_tokens, format_tokens = _tokens(Path(path))
    return Fingerprint(_hash(text_tokens), _hash(format_tokens))


def extract_text(path: Path) -> str:
    text_tokens, _ = _tokens(Path(path))
    return "\n".join(text_tokens)


def _tokens(path: Path) -> tuple[list, list]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx_tokens(path)
    if suffix == ".pptx":
        return _pptx_tokens(path)
    raise ValueError(f"지원하지 않는 형식입니다: {path.name}")


def _hash(tokens: list) -> str:
    data = json.dumps(tokens, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(data.encode("utf-8")).hexdigest()


def _coalesce(runs: Iterable[tuple[tuple, str]]) -> list[list]:
    """서식이 같은 이웃 런을 합친다. 편집기마다 런을 나누는 방식이 달라서다."""
    merged: list[list] = []
    for fmt, text in runs:
        if not text:
            continue
        if merged and merged[-1][0] == list(fmt):
            merged[-1][1] += text
        else:
            merged.append([list(fmt), text])
    return merged


# ---------------------------------------------------------------- docx


def _docx_tokens(path: Path) -> tuple[list, list]:
    document = docx.Document(str(path))
    text: list = []
    fmt: list = []

    def add_paragraph(p: Paragraph) -> None:
        text.append(p.text)
        style = p.style.name if p.style is not None else ""
        has_list = p._p.pPr is not None and p._p.pPr.numPr is not None
        align = str(p.alignment) if p.alignment is not None else ""
        fmt.append(["p", style, has_list, align, _coalesce(_docx_runs(p))])

    def walk(blocks: Iterable[Paragraph | Table]) -> None:
        for block in blocks:
            if isinstance(block, Paragraph):
                add_paragraph(block)
            else:
                fmt.append(["table", len(block.rows), len(block.columns)])
                for row in block.rows:
                    for cell in row.cells:
                        walk(_cell_blocks(cell))

    walk(document.iter_inner_content())
    for section in document.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:
                walk(part.paragraphs)

    images = sorted(
        hashlib.sha256(rel.target_part.blob).hexdigest()
        for rel in document.part.rels.values()
        if "image" in rel.reltype and not rel.is_external
    )
    fmt.append(["images", images])
    return text, fmt


def _cell_blocks(cell) -> Iterator[Paragraph | Table]:
    yield from cell.paragraphs
    yield from cell.tables


def _docx_runs(p: Paragraph) -> Iterator[tuple[tuple, str]]:
    for item in p.iter_inner_content():
        if isinstance(item, Hyperlink):
            for run in item.runs:
                yield (*_docx_run_format(run), item.address or ""), run.text
        else:
            yield (*_docx_run_format(item), ""), item.text


def _docx_run_format(run) -> tuple:
    font = run.font
    size = font.size.pt if font.size is not None else None
    return (
        bool(font.bold),
        bool(font.italic),
        bool(font.underline),
        bool(font.strike),
        size,
    )


# ---------------------------------------------------------------- pptx


def _pptx_tokens(path: Path) -> tuple[list, list]:
    presentation = pptx.Presentation(str(path))
    text: list = []
    fmt: list = []

    for index, slide in enumerate(presentation.slides):
        fmt.append(["slide", index])
        for shape in _iter_shapes(slide.shapes):
            fmt.append(
                ["shape", str(shape.shape_type), shape.left, shape.top, shape.width, shape.height]
            )
            if shape.has_text_frame:
                _pptx_text_frame(shape.text_frame, text, fmt)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        _pptx_text_frame(cell.text_frame, text, fmt)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                fmt.append(["image", hashlib.sha256(shape.image.blob).hexdigest()])
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None:
                text.append(notes.text)
                fmt.append(["notes", notes.text])
    return text, fmt


def _iter_shapes(shapes) -> Iterator:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _pptx_text_frame(frame, text: list, fmt: list) -> None:
    for p in frame.paragraphs:
        runs = []
        for run in p.runs:
            font = run.font
            size = font.size.pt if font.size is not None else None
            link = run.hyperlink.address or ""
            runs.append(
                ((bool(font.bold), bool(font.italic), bool(font.underline), size, link), run.text)
            )
        text.append("".join(r[1] for r in runs))
        fmt.append(
            ["p", p.level, str(p.alignment) if p.alignment is not None else "", _coalesce(runs)]
        )
