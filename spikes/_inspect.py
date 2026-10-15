"""내보낸 .docx/.pptx의 구성(글꼴, 표, 이미지 수 등)을 요약한다. V1 보고서용."""

from __future__ import annotations

from collections import Counter
from pathlib import Path

import docx
import pptx
from docx.oxml.ns import qn as w_qn
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn as a_qn


def describe(path: Path) -> dict:
    return _describe_docx(path) if path.suffix == ".docx" else _describe_pptx(path)


def _describe_docx(path: Path) -> dict:
    document = docx.Document(str(path))
    fonts: Counter[str] = Counter()
    for root in (document.element, document.styles.element):
        for rfonts in root.iter(w_qn("w:rFonts")):
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
                if name := rfonts.get(w_qn(attr)):
                    fonts[name] += 1
    paragraphs = document.paragraphs
    return {
        "문단": len(paragraphs),
        "제목 문단": sum(
            1
            for p in paragraphs
            if p.style is not None and p.style.name.lower().startswith(("heading", "title"))
        ),
        "목록 문단": sum(
            1 for p in paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None
        ),
        "표": len(document.tables),
        "이미지": sum(1 for r in document.part.rels.values() if "image" in r.reltype),
        "구역": len(document.sections),
        "사용 글꼴": ", ".join(name for name, _ in fonts.most_common()),
    }


def _describe_pptx(path: Path) -> dict:
    presentation = pptx.Presentation(str(path))
    kinds: Counter[str] = Counter()
    fonts: Counter[str] = Counter()
    notes = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                kinds["이미지"] += 1
            elif getattr(shape, "has_table", False) and shape.has_table:
                kinds["표"] += 1
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                kinds["차트"] += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                kinds["그룹"] += 1
            kinds["도형 전체"] += 1
        for tag in ("a:latin", "a:ea"):
            for el in slide.element.iter(a_qn(tag)):
                name = el.get("typeface")
                if name and not name.startswith("+"):
                    fonts[name] += 1
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes += bool(slide.notes_slide.notes_text_frame.text.strip())
    return {
        "슬라이드": len(presentation.slides),
        "슬라이드 크기(cm)": f"{presentation.slide_width / 360000:.1f} × "
        f"{presentation.slide_height / 360000:.1f}",
        **{k: kinds[k] for k in ("도형 전체", "이미지", "표", "차트", "그룹")},
        "발표자 노트 있는 슬라이드": notes,
        "사용 글꼴": ", ".join(name for name, _ in fonts.most_common()),
    }
