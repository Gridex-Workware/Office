import docx
import pptx
import pytest
from pptx.util import Inches

from officesync.content import extract_text, fingerprint


def make_docx(path, runs, heading=None):
    document = docx.Document()
    if heading:
        document.add_heading(heading, level=1)
    p = document.add_paragraph()
    for text, bold in runs:
        p.add_run(text).bold = bold
    document.save(path)
    return path


def make_pptx(path, text, left=None):
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(left or Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = text
    presentation.save(path)
    return path


def test_docx_resave_keeps_fingerprint(tmp_path):
    a = make_docx(tmp_path / "a.docx", [("안녕하세요 ", False), ("세계", True)], heading="제목")
    b = tmp_path / "b.docx"
    docx.Document(str(a)).save(b)
    assert fingerprint(a) == fingerprint(b)


def test_docx_split_runs_same_format_are_equal(tmp_path):
    a = make_docx(tmp_path / "a.docx", [("가나다", False)])
    b = make_docx(tmp_path / "b.docx", [("가나", False), ("다", False)])
    assert fingerprint(a) == fingerprint(b)


def test_docx_format_change_only(tmp_path):
    a = make_docx(tmp_path / "a.docx", [("가나다", False)])
    b = make_docx(tmp_path / "b.docx", [("가나다", True)])
    fa, fb = fingerprint(a), fingerprint(b)
    assert fa.text == fb.text
    assert fa.format != fb.format


def test_docx_text_change(tmp_path):
    a = make_docx(tmp_path / "a.docx", [("가나다", False)])
    b = make_docx(tmp_path / "b.docx", [("가나라", False)])
    fa, fb = fingerprint(a), fingerprint(b)
    assert fa.text != fb.text
    assert fa.format != fb.format


def test_docx_extract_text(tmp_path):
    a = make_docx(tmp_path / "a.docx", [("본문", False)], heading="제목")
    assert extract_text(a).splitlines() == ["제목", "본문"]


def test_pptx_resave_keeps_fingerprint(tmp_path):
    a = make_pptx(tmp_path / "a.pptx", "발표 내용")
    b = tmp_path / "b.pptx"
    pptx.Presentation(str(a)).save(b)
    assert fingerprint(a) == fingerprint(b)


def test_pptx_text_change(tmp_path):
    a = make_pptx(tmp_path / "a.pptx", "발표 내용")
    b = make_pptx(tmp_path / "b.pptx", "발표 내용 수정")
    assert fingerprint(a).text != fingerprint(b).text
    assert "발표 내용 수정" in extract_text(b)


def test_pptx_move_shape_is_format_change(tmp_path):
    a = make_pptx(tmp_path / "a.pptx", "발표 내용", left=Inches(1))
    b = make_pptx(tmp_path / "b.pptx", "발표 내용", left=Inches(2))
    fa, fb = fingerprint(a), fingerprint(b)
    assert fa.text == fb.text
    assert fa.format != fb.format


def test_unsupported_format(tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("x")
    with pytest.raises(ValueError):
        fingerprint(path)
